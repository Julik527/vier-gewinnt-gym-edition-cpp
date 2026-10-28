#!/usr/bin/env python3
"""Read-only preflight for the published JK AXIS Gym Edition artifacts.

Checks format validity, expected PDF page counts, every PDF page's rasterizability,
all presentation slides, official SVG assets, and local README links.
Does not modify or re-encode documents.
"""
from __future__ import annotations
import argparse
import hashlib
from pathlib import Path
import re
import sys
import xml.etree.ElementTree as ET

EXPECTED_PDFS = {
    "reference/Anforderungsanalyse_JK-CD-DE-2026-003.pdf": 19,
    "reference/Gesamt_Dokumentation_Vier_Gewinnt_Gym_Edition.pdf": 84,
    "reference/KI_Prompt_Dokumentation_JK-CD-DE-2026-003.pdf": 16,
    "reference/Meilenstein_2026-10-29_Pruefbericht.pdf": 7,
    "reference/UML_Klassendiagramm_JK-CD-DE-2026-003.pdf": 6,
    "praesentation/Vier_Gewinnt_Praesentation_FINAL.pdf": 17,
    "praesentation/Vier_Gewinnt_Vortragsleitfaden_FINAL.pdf": 2,
}
PPTX = "praesentation/Vier_Gewinnt_Praesentation_FINAL.pptx"
LEGACY_IMAGE_SHA_PREFIX = "bffb61e37fae"


def check(root: Path, presentation_only: bool = False) -> int:
    import fitz
    from pptx import Presentation

    errors: list[str] = []
    notes: list[str] = []
    slides = root / PPTX
    if not slides.is_file():
        errors.append(f"Missing presentation: {PPTX}")
    else:
        try:
            prs = Presentation(slides)
            if len(prs.slides) != 17:
                errors.append(f"Expected 17 slides, found {len(prs.slides)}")
            for index, slide in enumerate(prs.slides, start=1):
                for shape in slide.shapes:
                    if shape.has_text_frame and "JK-CD-DE-2026-003" in shape.text:
                        errors.append(f"Legacy footer in slide {index}")
                    if shape.shape_type == 13:
                        digest = hashlib.sha256(shape.image.blob).hexdigest()
                        if digest.startswith(LEGACY_IMAGE_SHA_PREFIX):
                            errors.append(f"Legacy raster logo in slide {index}")
            notes.append(f"PPTX: {len(prs.slides)} readable slides")
        except Exception as error:
            errors.append(f"PPTX cannot be opened: {error}")

    for path in [
        "assets/brand/jk-axis-precision-symbol.svg",
        "assets/brand/jk-axis-precision-symbol-inverse.svg",
    ]:
        file = root / path
        if presentation_only:
            break
        if not file.is_file():
            errors.append(f"Missing brand asset: {path}")
            continue
        try:
            svg = ET.parse(file).getroot()
            if svg.tag != "{http://www.w3.org/2000/svg}svg" or not svg.get("viewBox"):
                errors.append(f"Invalid SVG root/viewBox: {path}")
        except ET.ParseError as error:
            errors.append(f"SVG parsing error in {path}: {error}")

    if presentation_only:
        pdf_files = {PPTX.replace(".pptx", ".pdf"): 17}
    else:
        pdf_files = EXPECTED_PDFS

    for relative, expected_pages in pdf_files.items():
        file = root / relative
        if not file.is_file():
            errors.append(f"Missing PDF: {relative}")
            continue
        try:
            with fitz.open(file) as pdf:
                if pdf.is_encrypted:
                    errors.append(f"Unexpected encrypted PDF: {relative}")
                    continue
                if len(pdf) != expected_pages:
                    errors.append(f"{relative}: expected {expected_pages} pages, found {len(pdf)}")
                for index, page in enumerate(pdf, start=1):
                    if page.rect.width <= 0 or page.rect.height <= 0:
                        errors.append(f"{relative}: page {index} has invalid dimensions")
                        continue
                    # Rasterize all pages, not merely the first page.
                    pix = page.get_pixmap(matrix=fitz.Matrix(0.65, 0.65), alpha=False)
                    if pix.width <= 0 or pix.height <= 0:
                        errors.append(f"{relative}: page {index} rendered empty")
                notes.append(f"PDF: {relative} ({len(pdf)} pages rendered)")
        except Exception as error:
            errors.append(f"{relative}: cannot open/render PDF: {error}")

    if not presentation_only:
        readme = root / "README.md"
        if readme.is_file():
            md = readme.read_text(encoding="utf-8")
            for link in re.findall(r"\]\(([^)]+)\)", md):
                address = link.split("#", 1)[0]
                if not address or "://" in address or address.startswith("mailto:"):
                    continue
                if not (root / address).exists():
                    errors.append(f"Broken README link: {address}")
        else:
            errors.append("README.md is missing")

    print("\n".join(notes))
    if errors:
        print("\nQUALITY PREFLIGHT FAILED:\n- " + "\n- ".join(errors), file=sys.stderr)
        return 1
    print("QUALITY PREFLIGHT PASSED (technical files only; not a visual or school acceptance).")
    return 0


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--presentation-only", action="store_true")
    args = parser.parse_args()
    raise SystemExit(check(args.root, args.presentation_only))
